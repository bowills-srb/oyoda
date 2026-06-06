"""
MCP Data Collector Client

Python client for the STR Data Collector MCP server.
Used by the onboarding agent to scan local files for property data.

In production, this communicates with the MCP server via stdio.
For direct integration, we implement the same logic in Python.
"""

import asyncio
import json
import logging
import os
import re
from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path
from typing import Any, Dict, List, Optional, Tuple

logger = logging.getLogger(__name__)


# Supported file extensions
SUPPORTED_EXTENSIONS = {
    "spreadsheet": [".xlsx", ".xls", ".csv", ".tsv"],
    "document": [".docx", ".doc", ".pdf", ".txt", ".md", ".rtf"],
    "image": [".jpg", ".jpeg", ".png", ".gif", ".webp"],
    "data": [".json", ".xml", ".yaml", ".yml"],
}


# Property data patterns
PROPERTY_PATTERNS = {
    "wifi": re.compile(
        r"(?:wifi|wi-fi|wireless|network)\s*[:=]?\s*[\"']?([^\"'\n,]+)[\"']?",
        re.IGNORECASE
    ),
    "wifi_password": re.compile(
        r"(?:password|pwd|pass|key)\s*[:=]?\s*[\"']?([^\"'\n,]+)[\"']?",
        re.IGNORECASE
    ),
    "door_code": re.compile(
        r"(?:door|entry|access|lock|keypad)\s*(?:code)?\s*[:=]?\s*[\"']?(\d{4,6})[\"']?",
        re.IGNORECASE
    ),
    "gate_code": re.compile(
        r"(?:gate|community)\s*(?:code)?\s*[:=]?\s*[\"']?(\d{4,6})[\"']?",
        re.IGNORECASE
    ),
    "check_in": re.compile(
        r"(?:check[-\s]?in)\s*(?:time)?\s*[:=]?\s*[\"']?(\d{1,2}(?::\d{2})?\s*(?:am|pm)?)",
        re.IGNORECASE
    ),
    "check_out": re.compile(
        r"(?:check[-\s]?out)\s*(?:time)?\s*[:=]?\s*[\"']?(\d{1,2}(?::\d{2})?\s*(?:am|pm)?)",
        re.IGNORECASE
    ),
    "bedrooms": re.compile(
        r"(\d+)\s*(?:bed(?:room)?s?|br|bdrm)",
        re.IGNORECASE
    ),
    "bathrooms": re.compile(
        r"(\d+(?:\.\d)?)\s*(?:bath(?:room)?s?|ba)",
        re.IGNORECASE
    ),
    "sleeps": re.compile(
        r"(?:sleeps?|max\s*guests?|occupancy)\s*[:=]?\s*(\d+)",
        re.IGNORECASE
    ),
    "address": re.compile(
        r"(\d+\s+[A-Za-z\s]+(?:St|Street|Ave|Avenue|Rd|Road|Dr|Drive|Ln|Lane|Blvd|Way|Circle|Ct|Court)[^,\n]*)",
        re.IGNORECASE
    ),
    "property_code": re.compile(
        r"\b([A-Z]{2,}[-_]?\d{0,4}|[A-Z]{3,12})\b"
    ),
}


@dataclass
class FileAnalysis:
    """Result of analyzing a single file."""
    path: str
    file_name: str
    extension: str
    file_type: str
    size: int
    modified: datetime
    content_preview: Optional[str] = None
    structured_data: Optional[Any] = None
    extracted_data: Dict[str, List[Any]] = field(default_factory=dict)
    error: Optional[str] = None


@dataclass
class ScanResult:
    """Result of scanning a directory."""
    directory: str
    total_files: int
    spreadsheets: List[str] = field(default_factory=list)
    documents: List[str] = field(default_factory=list)
    data_files: List[str] = field(default_factory=list)
    images: List[str] = field(default_factory=list)


@dataclass
class PropertyData:
    """Extracted property data."""
    code: str
    operator_id: str
    sources: List[str] = field(default_factory=list)
    wifi_network: Optional[str] = None
    wifi_password: Optional[str] = None
    door_code: Optional[str] = None
    gate_code: Optional[str] = None
    check_in_time: Optional[str] = None
    check_out_time: Optional[str] = None
    bedrooms: Optional[int] = None
    bathrooms: Optional[float] = None
    sleeps: Optional[int] = None
    address: Optional[str] = None
    raw_data: Dict[str, Any] = field(default_factory=dict)


class DataCollector:
    """
    Collects property data from local files.
    
    Usage:
        collector = DataCollector()
        
        # Scan a directory
        scan = await collector.scan_directory("~/Documents/Properties")
        
        # Analyze files
        analyses = await collector.analyze_files(scan.spreadsheets + scan.documents)
        
        # Build property database
        properties = collector.build_property_database(analyses, "op_beach_habitats")
    """
    
    def __init__(self):
        self._xlsx = None
        self._pdf_parse = None
        self._mammoth = None
        self._load_parsers()
    
    def _load_parsers(self):
        """Load optional parsing libraries."""
        try:
            import openpyxl
            self._xlsx = openpyxl
        except ImportError:
            logger.warning("openpyxl not available - Excel parsing disabled")
        
        try:
            import PyPDF2
            self._pdf_parse = PyPDF2
        except ImportError:
            logger.warning("PyPDF2 not available - PDF parsing disabled")
        
        try:
            import docx
            self._docx = docx
        except ImportError:
            logger.warning("python-docx not available - Word parsing disabled")
    
    async def scan_directory(
        self,
        directory: str,
        recursive: bool = True,
        max_depth: int = 5,
    ) -> ScanResult:
        """
        Scan a directory for property-related files.
        
        Args:
            directory: Path to scan (supports ~ expansion)
            recursive: Whether to scan subdirectories
            max_depth: Maximum directory depth
        
        Returns:
            ScanResult with categorized file paths
        """
        # Expand ~ to home directory
        expanded_path = os.path.expanduser(directory)
        
        if not os.path.exists(expanded_path):
            raise ValueError(f"Directory not found: {expanded_path}")
        
        result = ScanResult(
            directory=expanded_path,
            total_files=0,
        )
        
        # Walk directory
        for root, dirs, files in os.walk(expanded_path):
            # Check depth
            depth = root[len(expanded_path):].count(os.sep)
            if depth >= max_depth:
                dirs.clear()  # Don't recurse deeper
                continue
            
            if not recursive and depth > 0:
                break
            
            # Skip common non-data directories
            dirs[:] = [d for d in dirs if d not in [
                "node_modules", ".git", "venv", "__pycache__",
                ".venv", "env", ".env"
            ]]
            
            for file in files:
                file_path = os.path.join(root, file)
                ext = os.path.splitext(file)[1].lower()
                
                if ext in SUPPORTED_EXTENSIONS["spreadsheet"]:
                    result.spreadsheets.append(file_path)
                elif ext in SUPPORTED_EXTENSIONS["document"]:
                    result.documents.append(file_path)
                elif ext in SUPPORTED_EXTENSIONS["data"]:
                    result.data_files.append(file_path)
                elif ext in SUPPORTED_EXTENSIONS["image"]:
                    result.images.append(file_path)
                
                result.total_files += 1
        
        logger.info(
            f"Scanned {expanded_path}: {len(result.spreadsheets)} spreadsheets, "
            f"{len(result.documents)} documents, {len(result.data_files)} data files"
        )
        
        return result
    
    async def analyze_file(self, file_path: str) -> FileAnalysis:
        """
        Analyze a single file and extract property data.
        
        Args:
            file_path: Path to the file
        
        Returns:
            FileAnalysis with extracted data
        """
        expanded_path = os.path.expanduser(file_path)
        
        if not os.path.exists(expanded_path):
            return FileAnalysis(
                path=expanded_path,
                file_name=os.path.basename(expanded_path),
                extension="",
                file_type="unknown",
                size=0,
                modified=datetime.now(),
                error=f"File not found: {expanded_path}",
            )
        
        stat = os.stat(expanded_path)
        file_name = os.path.basename(expanded_path)
        ext = os.path.splitext(file_name)[1].lower()
        
        analysis = FileAnalysis(
            path=expanded_path,
            file_name=file_name,
            extension=ext,
            file_type=self._get_file_type(ext),
            size=stat.st_size,
            modified=datetime.fromtimestamp(stat.st_mtime),
        )
        
        # Skip large files (> 10MB)
        if stat.st_size > 10 * 1024 * 1024:
            analysis.error = "File too large (>10MB)"
            return analysis
        
        try:
            content = await self._extract_content(expanded_path, ext)
            analysis.content_preview = content[:500] if content else None
            analysis.structured_data = await self._parse_structured(expanded_path, ext)
            
            # Extract property data
            if content:
                analysis.extracted_data = self._extract_property_data(content, file_name)
        except Exception as e:
            analysis.error = str(e)
            logger.error(f"Error analyzing {file_path}: {e}")
        
        return analysis
    
    async def analyze_files(
        self,
        file_paths: List[str],
        max_files: int = 100,
    ) -> List[FileAnalysis]:
        """
        Analyze multiple files.
        
        Args:
            file_paths: List of file paths
            max_files: Maximum files to process
        
        Returns:
            List of FileAnalysis results
        """
        results = []
        
        for path in file_paths[:max_files]:
            try:
                analysis = await self.analyze_file(path)
                results.append(analysis)
            except Exception as e:
                logger.error(f"Error analyzing {path}: {e}")
        
        return results
    
    def build_property_database(
        self,
        analyses: List[FileAnalysis],
        operator_id: str,
    ) -> List[PropertyData]:
        """
        Build a property database from analyzed files.
        
        Args:
            analyses: List of file analyses
            operator_id: Operator ID for the properties
        
        Returns:
            List of PropertyData objects
        """
        properties: Dict[str, PropertyData] = {}
        
        for analysis in analyses:
            if not analysis.extracted_data:
                continue
            
            data = analysis.extracted_data
            codes = data.get("property_codes", [])
            
            # If no property code found, try to use filename
            if not codes:
                # Try to extract from filename
                match = PROPERTY_PATTERNS["property_code"].search(analysis.file_name.upper())
                if match:
                    codes = [match.group(1)]
            
            for code in codes:
                if code not in properties:
                    properties[code] = PropertyData(
                        code=code,
                        operator_id=operator_id,
                    )
                
                prop = properties[code]
                prop.sources.append(analysis.path)
                
                # Merge data (first non-None wins)
                if not prop.wifi_network and data.get("wifi"):
                    prop.wifi_network = data["wifi"][0]
                if not prop.wifi_password and data.get("wifi_password"):
                    prop.wifi_password = data["wifi_password"][0]
                if not prop.door_code and data.get("door_code"):
                    prop.door_code = data["door_code"][0]
                if not prop.gate_code and data.get("gate_code"):
                    prop.gate_code = data["gate_code"][0]
                if not prop.check_in_time and data.get("check_in"):
                    prop.check_in_time = data["check_in"][0]
                if not prop.check_out_time and data.get("check_out"):
                    prop.check_out_time = data["check_out"][0]
                if not prop.bedrooms and data.get("bedrooms"):
                    prop.bedrooms = int(data["bedrooms"][0])
                if not prop.bathrooms and data.get("bathrooms"):
                    prop.bathrooms = float(data["bathrooms"][0])
                if not prop.sleeps and data.get("sleeps"):
                    prop.sleeps = int(data["sleeps"][0])
                if not prop.address and data.get("address"):
                    prop.address = data["address"][0]
        
        return list(properties.values())
    
    def _get_file_type(self, ext: str) -> str:
        """Get file type from extension."""
        for file_type, extensions in SUPPORTED_EXTENSIONS.items():
            if ext in extensions:
                return file_type
        return "unknown"
    
    async def _extract_content(self, file_path: str, ext: str) -> Optional[str]:
        """Extract text content from a file."""
        try:
            if ext in [".txt", ".md", ".csv", ".tsv"]:
                with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
                    return f.read()
            
            elif ext == ".json":
                with open(file_path, "r", encoding="utf-8") as f:
                    return f.read()
            
            elif ext == ".docx" and hasattr(self, "_docx"):
                doc = self._docx.Document(file_path)
                return "\n".join([p.text for p in doc.paragraphs])
            
            elif ext == ".pdf" and self._pdf_parse:
                with open(file_path, "rb") as f:
                    reader = self._pdf_parse.PdfReader(f)
                    text = ""
                    for page in reader.pages:
                        text += page.extract_text() or ""
                    return text
            
            elif ext in [".xlsx", ".xls"] and self._xlsx:
                wb = self._xlsx.load_workbook(file_path, data_only=True)
                text_parts = []
                for sheet in wb.worksheets:
                    for row in sheet.iter_rows(values_only=True):
                        row_text = " ".join(str(cell) for cell in row if cell)
                        if row_text.strip():
                            text_parts.append(row_text)
                return "\n".join(text_parts)
        
        except Exception as e:
            logger.error(f"Error extracting content from {file_path}: {e}")
        
        return None
    
    async def _parse_structured(self, file_path: str, ext: str) -> Optional[Any]:
        """Parse structured data from a file."""
        try:
            if ext == ".json":
                with open(file_path, "r") as f:
                    return json.load(f)
            
            elif ext in [".xlsx", ".xls"] and self._xlsx:
                wb = self._xlsx.load_workbook(file_path, data_only=True)
                sheets = {}
                for sheet in wb.worksheets:
                    rows = []
                    headers = None
                    for i, row in enumerate(sheet.iter_rows(values_only=True)):
                        if i == 0:
                            headers = [str(cell) if cell else f"col_{j}" for j, cell in enumerate(row)]
                        else:
                            if headers:
                                row_dict = {headers[j]: cell for j, cell in enumerate(row) if j < len(headers)}
                                rows.append(row_dict)
                    sheets[sheet.title] = rows
                return sheets
            
            elif ext == ".csv":
                import csv
                with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
                    reader = csv.DictReader(f)
                    return list(reader)
        
        except Exception as e:
            logger.error(f"Error parsing structured data from {file_path}: {e}")
        
        return None
    
    def _extract_property_data(
        self,
        text: str,
        file_name: str = "",
    ) -> Dict[str, List[Any]]:
        """
        Extract property data from text content.
        
        Args:
            text: Text content to analyze
            file_name: Optional filename for additional context
        
        Returns:
            Dictionary of extracted data by type
        """
        data = {
            "property_codes": [],
            "wifi": [],
            "wifi_password": [],
            "door_code": [],
            "gate_code": [],
            "check_in": [],
            "check_out": [],
            "bedrooms": [],
            "bathrooms": [],
            "sleeps": [],
            "address": [],
        }
        
        # Try to get property code from filename
        file_match = PROPERTY_PATTERNS["property_code"].search(file_name.upper())
        if file_match:
            data["property_codes"].append(file_match.group(1))
        
        # Extract patterns from text
        for key, pattern in PROPERTY_PATTERNS.items():
            if key == "property_code":
                # Only add property codes that look valid (not common words)
                matches = pattern.findall(text)
                for match in matches:
                    if len(match) >= 4 and match.upper() not in ["WITH", "FROM", "THAT", "THIS", "HAVE", "BEEN"]:
                        if match not in data["property_codes"]:
                            data["property_codes"].append(match)
            else:
                matches = pattern.findall(text)
                for match in matches:
                    cleaned = match.strip() if isinstance(match, str) else match
                    if cleaned and cleaned not in data[key]:
                        data[key].append(cleaned)
        
        return data


# Singleton instance
_collector: Optional[DataCollector] = None


def get_data_collector() -> DataCollector:
    """Get or create data collector instance."""
    global _collector
    if _collector is None:
        _collector = DataCollector()
    return _collector
